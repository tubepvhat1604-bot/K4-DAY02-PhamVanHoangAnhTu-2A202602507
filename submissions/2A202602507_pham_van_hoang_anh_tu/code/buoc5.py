"""buoc5.py - Bước 5: gom mọi log thật thành results.xlsx (GUIDE.md mục 6.1), vẽ biểu đồ tổng hợp cho báo cáo,
ảnh lỗi Chinee apple <-> Snake weed, và tự kiểm tra bài nộp. Không huấn luyện, không chạy lại test.

Nguồn số liệu (không gõ tay con số nào):
  <runs_dir>/<exp_id>/seed<k>/done.json         -> Backbones, Training, Final (val của T00)
  <sub_dir>/logs/inference_val.csv, latency.csv -> Inference, Latency (Bước 3)
  <sub_dir>/logs/eval/<tag>_per_seed.csv, _per_class.csv (eval.py score) -> Final, PerClass (Bước 4)
  <sub_dir>/predictions/F01_seed<k>_val.csv      -> macro-F1 val của chung kết (sau TS)
"""
from __future__ import annotations

import json
import re
from pathlib import Path

import numpy as np
import pandas as pd

# Trục (GUIDE mục 3) và điểm khác T00 của từng thí nghiệm Bước 2
TRAIN_META = {
    "T00": ("-", "công thức nền (tinh chỉnh toàn bộ, aug cơ bản, CE, không sampler, không EMA)"),
    "T01": ("A. Khởi tạo", "đóng băng backbone, chỉ train head"),
    "T02": ("A. Khởi tạo", "train từ đầu (không trọng số tiền huấn luyện)"),
    "T03": ("B. Augmentation", "+ TrivialAugment"),
    "T04": ("B. Augmentation", "+ CutMix (trộn cả nhãn)"),
    "T05": ("B. Augmentation", "+ lật dọc và xoay 90° (dihedral)"),
    "T06": ("C. Loss", "CE + label smoothing 0,1"),
    "T07": ("C. Loss", "focal loss γ = 2"),
    "T08": ("C. Loss", "CE trọng số class-balanced (β = 0,9999)"),
    "T09": ("D. Cân bằng mẫu", "sampler cân bằng lớp"),
    "T10": ("F. Chính quy hoá", "EMA trọng số 0,999"),
    "T11": ("B + D (kết hợp)", "TrivialAugment + sampler cân bằng"),
}
BACKBONE_NOTE = {
    "B01": "ResNet", "B02": "ConvNeXt (chọn đi tiếp)", "B03": "transformer (DeiT)",
    "B04": "mạng nhẹ", "B05": "mạng nhẹ", "B06": "transformer (Swin)",
}
SHORT = {"T01": "đóng băng", "T02": "từ đầu", "T03": "TrivialAug", "T04": "CutMix", "T05": "dihedral",
         "T06": "LS 0,1", "T07": "focal γ2", "T08": "CB weight", "T09": "sampler CB", "T10": "EMA",
         "T11": "TrivialAug\n+ sampler"}


def _done(runs_dir, prefix):
    rows = []
    for f in sorted(Path(runs_dir).glob(f"{prefix}*/seed*/done.json")):
        d = json.loads(f.read_text())
        pc = d.get("val_f1_per_class") or [np.nan] * 9
        d.update(f1_chinee=pc[0], f1_snake=pc[7], f1_negative=pc[8])
        rows.append(d)
    if not rows:
        raise FileNotFoundError(f"Không thấy {runs_dir}/{prefix}*/seed*/done.json")
    return pd.DataFrame(rows).sort_values(["exp_id", "seed"]).reset_index(drop=True)


def _pm(s):
    s = pd.Series(s, dtype=float).dropna()
    return f"{s.mean():.4f} ± {s.std(ddof=1):.4f}" if len(s) > 1 else f"{s.mean():.4f}"


def sheet_backbones(runs_dir):
    d = _done(runs_dir, "B")
    return pd.DataFrame({
        "exp_id": d.exp_id, "backbone": d.backbone, "tag trọng số (timm)": d.weight_tag,
        "#tham số (M)": d.params_M, "GMAC": d.gmacs, "độ phân giải (px)": d.img_size, "epoch": d.epochs,
        "seed": d.seed, "best epoch": d.best_epoch, "macro-F1 val": d.val_macro_f1, "top-1 val": d.val_top1,
        "thời gian train/epoch (s)": d.train_time_per_epoch_s,
        "độ trễ batch-1 sơ bộ (ms, AMP, 30 lần)": d.latency_b1_ms_prelim,
        "ghi chú": d.exp_id.map(BACKBONE_NOTE).fillna("") + "; Bước 1 chạy trên Colab",
    })


def sheet_training(runs_dir):
    d = _done(runs_dir, "T")
    base = d.loc[d.exp_id == "T00", "val_macro_f1"]
    noise = base.std(ddof=1)
    delta = d.val_macro_f1 - base.mean()
    note = []
    for e, dl in zip(d.exp_id, delta):
        if e == "T00":
            note.append(f"mốc 3 seed: {_pm(base)}")
        elif abs(dl) > 2 * noise:
            note.append("khác rõ (|Δ| > 2 std của T00)")
        else:
            note.append("không phân biệt được (|Δ| ≤ 2 std, 1 seed)")
    return pd.DataFrame({
        "exp_id": d.exp_id, "backbone": d.backbone, "trục thay đổi (A–G)": d.exp_id.map(lambda e: TRAIN_META[e][0]),
        "khác T00 ở điểm nào": d.exp_id.map(lambda e: TRAIN_META[e][1]), "seed": d.seed,
        "macro-F1 val": d.val_macro_f1, "top-1 val": d.val_top1, "Δ macro-F1 so với T00 (mean 3 seed)": delta,
        "ECE val": d.val_ece, "F1 Chinee apple val": d.f1_chinee, "F1 Snake weed val": d.f1_snake,
        "F1 Negative val": d.f1_negative, "thời gian train/epoch (s)": d.train_time_per_epoch_s,
        "ghi chú": note,
    })


def sheet_inference(sub_dir):
    r = pd.read_csv(Path(sub_dir) / "logs" / "inference_val.csv")
    lat = pd.read_csv(Path(sub_dir) / "logs" / "latency.csv")
    b32 = {row.label.replace(" b32", ""): row.images_per_s for row in lat[lat.batch == 32].itertuples()}
    note = []
    for row in r.itertuples():
        n = []
        if row.method.startswith("I07"):
            n.append(f"T = {row.T:.3f}; ECE {row.val_ece_before:.4f} → {row.val_ece:.4f} (cross-fit {row.val_ece_crossfit:.4f})")
        if row.method == "I08_fp16":
            n.append(f"lệch logit tối đa so với FP32 = {row.max_abs_logit_diff_vs_fp32:.4f}")
        if row.method.startswith("I02"):
            n.append("5 crop 224 từ ảnh 256")
        note.append("; ".join(n))
    return pd.DataFrame({
        "exp_id": r.method, "phương pháp": r.method.str.split("_", n=1).str[1],
        "mô hình/checkpoint": r.model, "K (số view hoặc số mô hình)": r.views,
        "gộp": r.aggregate, "độ phân giải vào (px)": r.input_size,
        "macro-F1 val (F01 seed 0)": r.val_macro_f1, "macro-F1 val mean 3 seed": r.val_macro_f1_mean_3seed,
        "std 3 seed": r.val_macro_f1_std_3seed, "top-1 val": r.val_top1, "ECE val": r.val_ece,
        "Δ macro-F1 so với I00": r.delta_vs_I00,
        "p50 b1 (ms)": r.lat_b1_p50_ms, "p95 b1 (ms)": r.lat_b1_p95_ms, "p99 b1 (ms)": r.lat_b1_p99_ms,
        "thông lượng b1 (ảnh/s)": 1000.0 / r.lat_b1_p50_ms,
        "thông lượng b32 (ảnh/s)": r.method.map(lambda m: b32.get(m, np.nan)),
        "chi phí tương đối so với I00 (p50)": r.cost_vs_I00, "ghi chú": note,
    })


def sheet_latency(sub_dir):
    l = pd.read_csv(Path(sub_dir) / "logs" / "latency.csv")
    return pd.DataFrame({
        "cấu hình": l.label, "GPU": l.gpu, "dtype": l.dtype, "batch": l.batch, "độ phân giải (px)": l.img_size,
        "gộp BN": l.fused_bn.map({True: "có", False: "không"}), "p50 (ms)": l.p50, "p95 (ms)": l.p95,
        "p99 (ms)": l.p99, "ảnh/s": l.images_per_s, "số lần đo": l.n, "warmup": l.warmup,
        "tiền xử lý": l.preprocessing, "torch": l.torch,
    })


def _val_f1_from_preds(sub_dir, final_id, seeds):
    import eval as ev
    out = {}
    for k in seeds:
        p = ev.read_pred(str(Path(sub_dir) / "predictions" / f"{final_id}_seed{k}_val.csv"))
        out[k] = ev.compute_metrics(p.y_true, p.y_pred, p.probs)["macro_f1"]
    return out


def sheet_final(runs_dir, sub_dir, final_id="F01", base_id="T00"):
    ev_dir = Path(sub_dir) / "logs" / "eval"
    choice = json.loads((Path(sub_dir) / "logs" / "buoc3_choice.json").read_text())
    temps = json.loads((Path(sub_dir) / "logs" / "buoc4_test_info.json").read_text())["temperature"]
    fin = pd.read_csv(ev_dir / f"{final_id}_per_seed.csv")
    unc = pd.read_csv(ev_dir / f"{final_id}_uncal_per_seed.csv").set_index("seed")
    bas = pd.read_csv(ev_dir / f"{base_id}_per_seed.csv")
    val_f = _val_f1_from_preds(sub_dir, final_id, fin.seed)
    tr = _done(runs_dir, base_id).set_index("seed")
    cfg_f = f"ConvNeXt-T in12k_ft_in1k + T00 + TrivialAugment (= T03), 12 epoch; suy luận {choice['method']} + temperature scaling (T khớp trên val)"
    cfg_b = "ConvNeXt-T in12k_ft_in1k + T00 (công thức nền); suy luận I00 1 view 224, không TS"
    rows = []
    for r in fin.itertuples():
        rows.append({"exp_id": final_id, "cấu hình": cfg_f, "seed": r.seed, "macro-F1 val": val_f[r.seed],
                     "macro-F1 test": r.macro_f1, "top-1 test": r.top1, "balanced acc test": r.balanced_acc,
                     "ECE test": r.ece, "ECE test trước TS": unc.loc[r.seed, "ece"], "T": temps[f"seed{r.seed}"]})
    for r in bas.itertuples():
        rows.append({"exp_id": f"{base_id} + I00 (mốc)", "cấu hình": cfg_b, "seed": r.seed,
                     "macro-F1 val": tr.loc[r.seed, "val_macro_f1"], "macro-F1 test": r.macro_f1, "top-1 test": r.top1,
                     "balanced acc test": r.balanced_acc, "ECE test": r.ece, "ECE test trước TS": np.nan, "T": np.nan})
    df = pd.DataFrame(rows)
    out = [df]
    for name, g in df.groupby("exp_id", sort=False):
        agg = {"exp_id": name, "cấu hình": "mean ± std qua 3 seed (std mẫu, ddof = 1)", "seed": "mean ± std"}
        for c in ["macro-F1 val", "macro-F1 test", "top-1 test", "balanced acc test", "ECE test", "ECE test trước TS", "T"]:
            agg[c] = _pm(g[c]) if g[c].notna().any() else ""
        out.append(pd.DataFrame([agg]))
    return pd.concat(out, ignore_index=True)


def sheet_perclass(sub_dir, final_id="F01", base_id="T00"):
    ev_dir = Path(sub_dir) / "logs" / "eval"
    rows = []
    for tag, label in ((final_id, f"{final_id} (chung kết = tốt nhất)"), (base_id, f"{base_id} + I00 (mốc)")):
        pc = pd.read_csv(ev_dir / f"{tag}_per_class.csv")
        for r in pc.itertuples():
            rows.append({"cấu hình": label, "lớp": r[1], "số ảnh test": r.support,
                         "precision": _pm_ms(r.precision_mean, r.precision_std),
                         "recall": _pm_ms(r.recall_mean, r.recall_std), "F1": _pm_ms(r.f1_mean, r.f1_std),
                         "precision mean": r.precision_mean, "recall mean": r.recall_mean, "F1 mean": r.f1_mean})
    return pd.DataFrame(rows)


def _pm_ms(m, s):
    return f"{m:.4f} ± {s:.4f}"


def sheet_summary(bb, tr, inf, fin, final_id="F01"):
    """Dòng chung kết + mốc, rồi top 10 theo macro-F1 val (mỗi exp_id một dòng; nhiều seed thì lấy mean)."""
    p95 = inf.set_index("exp_id")["p95 b1 (ms)"]
    cand = []
    for _, r in bb.iterrows():
        cand.append((r["exp_id"], "Backbones", r["backbone"], r["macro-F1 val"],
                     r["độ trễ batch-1 sơ bộ (ms, AMP, 30 lần)"], "1 seed; độ trễ sơ bộ (30 lần, Colab)"))
    for e, g in tr.groupby("exp_id"):
        cand.append((e, "Training", TRAIN_META[e][1], g["macro-F1 val"].mean(), p95["I00_1view"],
                     f"{len(g)} seed" + (", mean" if len(g) > 1 else "") + "; suy luận I00"))
    for _, r in inf.iterrows():
        m3 = r["macro-F1 val mean 3 seed"]
        cand.append((r["exp_id"], "Inference", f"F01 + {r['phương pháp']}",
                     m3 if pd.notna(m3) else r["macro-F1 val (F01 seed 0)"], r["p95 b1 (ms)"],
                     "mean 3 seed" if pd.notna(m3) else
                     ("3 checkpoint F01 seed 0-2" if r["exp_id"][:3] in ("I05", "I06") else "F01 seed 0")))
    df = pd.DataFrame(cand, columns=["exp_id", "sheet", "mô tả", "macro-F1 val", "p95 b1 (ms)", "ghi chú"])
    top = df.sort_values("macro-F1 val", ascending=False).head(10).copy()
    top.insert(0, "hạng", [str(k) for k in range(1, len(top) + 1)])
    head = []
    for _, r in fin[fin["seed"].astype(str).str.contains("mean")].iterrows():
        is_final = r["exp_id"] == final_id
        head.append({"hạng": "CHUNG KẾT" if is_final else "MỐC", "exp_id": r["exp_id"], "sheet": "Final",
                     "mô tả": "ConvNeXt-T + TrivialAugment + res288 + TS" if is_final else "ConvNeXt-T + T00 + I00",
                     "macro-F1 val": r["macro-F1 val"], "macro-F1 test": r["macro-F1 test"],
                     "top-1 test": r["top-1 test"], "ECE test": r["ECE test"],
                     "p95 b1 (ms)": p95["I04_res288" if is_final else "I00_1view"], "ghi chú": "mean ± std 3 seed"})
    out = pd.concat([pd.DataFrame(head), top], ignore_index=True)
    out["chi phí so với I00 (p95)"] = out["p95 b1 (ms)"] / p95["I00_1view"]
    cols = ["hạng", "exp_id", "sheet", "mô tả", "macro-F1 val", "macro-F1 test", "top-1 test", "ECE test",
            "p95 b1 (ms)", "chi phí so với I00 (p95)", "ghi chú"]
    return out[cols]


def _format(path, highlight):
    from openpyxl import load_workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter
    wb = load_workbook(path)
    fill = PatternFill("solid", fgColor="FFF2CC")
    for ws in wb.worksheets:
        ws.freeze_panes = "A2"
        for c in ws[1]:
            c.font = Font(bold=True)
            c.alignment = Alignment(wrap_text=True, vertical="top")
        for col in ws.iter_cols(min_row=2):
            for c in col:
                if isinstance(c.value, float):
                    c.number_format = "0.0000"
        for i, col in enumerate(ws.iter_cols(min_row=1, max_row=min(ws.max_row, 40)), 1):
            w = max(len(str(c.value)) if c.value is not None else 0 for c in col)
            ws.column_dimensions[get_column_letter(i)].width = min(max(10, w + 2), 60)
        rule = highlight.get(ws.title)
        if rule:
            for row in ws.iter_rows(min_row=2):
                if rule(row):
                    for c in row:
                        c.fill = fill
                        c.font = Font(bold=True)
    wb.save(path)


def build_results(runs_dir, sub_dir, out=None, final_id="F01", base_id="T00"):
    """Ghi <sub_dir>/results.xlsx từ log thật và trả về dict các bảng."""
    bb, tr = sheet_backbones(runs_dir), sheet_training(runs_dir)
    inf, lat = sheet_inference(sub_dir), sheet_latency(sub_dir)
    fin, pc = sheet_final(runs_dir, sub_dir, final_id, base_id), sheet_perclass(sub_dir, final_id, base_id)
    sm = sheet_summary(bb, tr, inf, fin, final_id)
    sheets = {"Summary": sm, "Backbones": bb, "Training": tr, "Inference": inf, "Final": fin, "PerClass": pc,
              "Latency": lat}
    out = Path(out or Path(sub_dir) / "results.xlsx")
    with pd.ExcelWriter(out, engine="openpyxl") as w:
        for name, df in sheets.items():
            df.to_excel(w, sheet_name=name, index=False)
    best_bb = bb.loc[bb["macro-F1 val"].idxmax(), "exp_id"]
    tr_mean = tr.groupby("exp_id")["macro-F1 val"].mean()
    best_tr = tr_mean.drop([e for e in ("T11",) if e in tr_mean]).idxmax()
    best_inf = inf.loc[inf["macro-F1 val mean 3 seed"].idxmax(), "exp_id"]
    _format(out, {
        "Summary": lambda r: r[1].value == final_id,
        "Backbones": lambda r: r[0].value == best_bb,
        "Training": lambda r: r[0].value == best_tr,
        "Inference": lambda r: r[0].value == best_inf,
        "Final": lambda r: r[0].value == final_id and "mean" in str(r[2].value),
        "PerClass": lambda r: r[1].value in ("Chinee apple", "Snake weed"),
        "Latency": lambda r: r[0].value == best_inf,
    })
    return sheets


# --------------------------------------------------------------------------- #
# Biểu đồ tổng hợp cho báo cáo
# --------------------------------------------------------------------------- #
def plot_backbones(bb, path):
    import matplotlib
    matplotlib.use("Agg")
    import matplotlib.pyplot as plt
    fig, ax = plt.subplots(1, 2, figsize=(12, 4.6), dpi=130)
    for a, x, xl in ((ax[0], "độ trễ batch-1 sơ bộ (ms, AMP, 30 lần)", "độ trễ batch 1 (ms, đo sơ bộ 30 lần ở Bước 1)"),
                     (ax[1], "#tham số (M)", "số tham số (M)")):
        a.scatter(bb[x], bb["macro-F1 val"], s=40 + 3 * bb["GMAC"] * 10, alpha=.7)
        for _, r in bb.iterrows():
            a.annotate(f"{r['exp_id']} {r['backbone'].split('_patch')[0]}", (r[x], r["macro-F1 val"]),
                       fontsize=7, xytext=(5, -10 if r["exp_id"] == "B06" and x == "#tham số (M)" else 4),
                       textcoords="offset points")
        a.margins(x=0.2)
        a.set_xlabel(xl); a.set_ylabel("macro-F1 val (seed 0)"); a.grid(alpha=.3)
    fig.suptitle("Bước 1: macro-F1 val theo độ trễ và số tham số (cỡ điểm ~ GMAC)")
    fig.tight_layout(); fig.savefig(path); plt.close(fig)


def plot_ablation(tr, path):
    import matplotlib
    matplotlib.use("Agg")
    import matplotlib.pyplot as plt
    base = tr.loc[tr.exp_id == "T00", "macro-F1 val"]
    m, s = base.mean(), base.std(ddof=1)
    g = tr[tr.exp_id != "T00"].groupby("exp_id")["macro-F1 val"].mean()
    g = g[g > 0.9]                                      # T01, T02 kém rất xa: ghi trong chú thích
    fig, ax = plt.subplots(figsize=(9, 4.2), dpi=130)
    ax.axhspan(m - s, m + s, color="gray", alpha=.2, label="T00 mean ± 1 std (3 seed)")
    ax.axhspan(m - 2 * s, m + 2 * s, color="gray", alpha=.08, label="T00 mean ± 2 std")
    ax.axhline(m, color="gray", ls="--", lw=1)
    ax.scatter(np.zeros(len(base)) - 0.5, base, color="black", s=18, label="T00 từng seed")
    ax.bar(range(1, len(g) + 1), g.values - m, bottom=m, color=["tab:green" if v > m else "tab:red" for v in g.values])
    ax.set_xticks([-0.5] + list(range(1, len(g) + 1)), ["T00"] + [f"{e}\n{SHORT[e]}" for e in g.index], fontsize=7)
    lo, hi = min(g.min(), base.min()), max(g.max(), base.max())
    ax.set_ylim(lo - 0.004, hi + 0.004)
    dropped = tr[(tr.exp_id != "T00") & (tr["macro-F1 val"] <= 0.9)]
    ax.set_title("Bước 2: macro-F1 val so với T00 (1 seed mỗi thí nghiệm)" +
                 ("\nkhông vẽ: " + ", ".join(f"{e} = {v:.4f}" for e, v in zip(dropped.exp_id, dropped["macro-F1 val"]))
                  if len(dropped) else ""),
                 fontsize=9)
    ax.set_ylabel("macro-F1 val"); ax.legend(fontsize=8, loc="lower right"); ax.grid(alpha=.3, axis="y")
    fig.tight_layout(); fig.savefig(path); plt.close(fig)


def plot_errors(pred_csv, images_dir, path, pairs=((0, 7), (7, 0)), n=6, class_names=None):
    """Lưới ảnh test bị đoán sai cho từng cặp (nhãn thật, dự đoán), kèm xác suất. Chỉ để phân tích lỗi sau Bước 4."""
    import matplotlib
    matplotlib.use("Agg")
    import matplotlib.pyplot as plt
    from PIL import Image
    if class_names is None:
        from dataset import CLASS_NAMES as class_names
    d = pd.read_csv(pred_csv)
    fig, axes = plt.subplots(len(pairs), n, figsize=(2.2 * n, 2.6 * len(pairs)), dpi=110, squeeze=False)
    counts = {}
    for i, (t, p) in enumerate(pairs):
        sel = d[(d.y_true == t) & (d.y_pred == p)].sort_values(f"p{p}", ascending=False)
        counts[(t, p)] = len(sel)
        for j in range(n):
            a = axes[i, j]
            a.axis("off")
            if j < len(sel):
                r = sel.iloc[j]
                a.imshow(Image.open(Path(images_dir) / r.Filename).convert("RGB"))
                a.set_title(f"{r.Filename[:15]}\np({class_names[p][:8]})={r[f'p{p}']:.2f} p(đúng)={r[f'p{t}']:.2f}",
                            fontsize=6)
        axes[i, 0].text(-0.08, 0.5, f"thật: {class_names[t]}\nđoán: {class_names[p]}\n({len(sel)} ảnh)",
                        transform=axes[i, 0].transAxes, ha="right", va="center", fontsize=8)
    fig.suptitle(f"Ảnh test bị đoán sai ({Path(pred_csv).name})", fontsize=9)
    fig.tight_layout(); fig.savefig(path, bbox_inches="tight"); plt.close(fig)
    return counts


def top_confusions(confusion_csv, class_names=None, k=6):
    """Các cặp nhầm nhiều nhất (cộng 3 seed) từ <tag>_confusion_sum.csv của eval.py."""
    cm = pd.read_csv(confusion_csv, index_col=0).to_numpy()
    if class_names is None:
        from dataset import CLASS_NAMES as class_names
    off = [(cm[i, j], class_names[i], class_names[j]) for i in range(len(cm)) for j in range(len(cm)) if i != j]
    return pd.DataFrame(sorted(off, reverse=True)[:k], columns=["số ảnh (3 seed)", "nhãn thật", "dự đoán"])


# --------------------------------------------------------------------------- #
# Tự kiểm tra bài nộp
# --------------------------------------------------------------------------- #
def check_submission(sub_dir, final_id="F01", base_id="T00", seeds=(0, 1, 2)):
    """Trả về list vấn đề (rỗng = đạt): đủ sản phẩm, mỗi exp_id huấn luyện có ảnh curves/, đủ file dự đoán."""
    sub = Path(sub_dir)
    probs = []
    for p in ("results.xlsx", "report.md", "README.md", "code", "curves", "predictions"):
        if not (sub / p).exists():
            probs.append(f"thiếu {p}")
    if (sub / "results.xlsx").exists():
        xl = pd.read_excel(sub / "results.xlsx", sheet_name=None)
        need = {"Backbones", "Training", "Inference", "Final", "PerClass", "Latency", "Summary"}
        probs += [f"results.xlsx thiếu sheet {s}" for s in need - set(xl)]
        ids = set(xl.get("Backbones", pd.DataFrame(columns=["exp_id"])).exp_id) | \
            set(xl.get("Training", pd.DataFrame(columns=["exp_id"])).exp_id) | {final_id}
        pngs = [f.name for f in (sub / "curves").glob("*.png")]
        for e in sorted(ids):
            if not any(re.match(rf"{e}_", f) for f in pngs):
                probs.append(f"curves/ thiếu ảnh cho {e}")
    for k in seeds:
        for name in (f"{final_id}_seed{k}_test.csv", f"{base_id}_seed{k}_test.csv"):
            if not (sub / "predictions" / name).exists():
                probs.append(f"predictions/ thiếu {name}")
    big = [f for f in sub.rglob("*") if f.is_file() and f.stat().st_size > 20e6]
    probs += [f"file quá lớn (không nên commit): {f.relative_to(sub)}" for f in big]
    probs += [f"không được commit checkpoint: {f.relative_to(sub)}" for f in sub.rglob("*.pt")]
    return probs
